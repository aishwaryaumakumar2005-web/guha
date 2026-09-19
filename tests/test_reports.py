"""Batch 1 reports-module bug fixes (B1-B5) and Batch 2 functionality (F1-F5)."""
from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models import (Account, Course, Student, FeeRecord, Expense, Company,
                        ExpenseCategory, student_courses)
from app.routes.reports import (course_wise_income_summary,
                                filter_by_company_methods,
                                payment_method_period_breakdown)


def _mk_company(app, name):
    with app.app_context():
        c = Company(name=name, code=name.replace(' ', '').upper()[:6])
        db.session.add(c)
        db.session.commit()
        return c.id


def _mk_course(app, code, company_id=None):
    with app.app_context():
        c = Course(name='Course ' + code, code=code, description='',
                   duration_weeks=8, duration_unit='weeks', fees=1000.0,
                   gst_applicable=False, company_id=company_id)
        db.session.add(c)
        db.session.commit()
        return c.id


def _mk_student(app, seq, course_ids):
    with app.app_context():
        s = Student(name=f'Student {seq}', email=f'student{seq}@reports.guha.test',
                    phone='9' * 10, status='Active')
        db.session.add(s)
        db.session.flush()
        for cid in course_ids:
            s.courses.append(Course.query.get(cid))
        db.session.commit()
        return s.id


def _set_enrollment(app, student_id, course_id, status=None, enrolled_on=None, completed_on=None):
    with app.app_context():
        vals = {}
        if status is not None:
            vals['status'] = status
        if enrolled_on is not None:
            vals['enrolled_on'] = enrolled_on
        if completed_on is not None:
            vals['completed_on'] = completed_on
        db.session.execute(
            student_courses.update().where(
                student_courses.c.student_id == student_id,
                student_courses.c.course_id == course_id,
            ).values(**vals))
        db.session.commit()


def _add_fee(app, student_id, amount, on, method='Cash', company_id=None):
    with app.app_context():
        db.session.add(FeeRecord(student_id=student_id, amount_paid=amount,
                                 taxable_amount=amount, gst_amount=0.0,
                                 payment_date=on, payment_method=method,
                                 company_id=company_id))
        db.session.commit()


def _add_expense(app, amount, on, company_id=None, method='Cash'):
    with app.app_context():
        cat = ExpenseCategory.query.filter_by(name='Rent').first()
        db.session.add(Expense(amount=amount, expense_date=on, category_id=cat.id,
                               description='scratch', payment_method=method,
                               company_id=company_id))
        db.session.commit()


# -- B1: course-wise income ignores dropped/inactive/future enrollments --------
def test_course_wise_income_excludes_dropped_course(app):
    ca = _mk_course(app, 'CWA')
    cb = _mk_course(app, 'CWB')
    sid = _mk_student(app, 1, [ca, cb])
    _set_enrollment(app, sid, ca, status='Enrolled', enrolled_on=date(2026, 9, 1))
    _set_enrollment(app, sid, cb, status='Dropped', enrolled_on=date(2026, 8, 1),
                    completed_on=date(2026, 8, 15))
    _add_fee(app, sid, 3000.0, date(2026, 9, 10))
    with app.app_context():
        m = course_wise_income_summary(date(2026, 9, 1), date(2026, 9, 30))
    assert m.get(ca) == 3000.0
    assert cb not in m


def test_course_wise_income_excludes_future_enrollment(app):
    ca = _mk_course(app, 'FWA')
    cb = _mk_course(app, 'FWB')
    sid = _mk_student(app, 2, [ca, cb])
    _set_enrollment(app, sid, ca, status='Enrolled', enrolled_on=date(2026, 9, 20))
    _set_enrollment(app, sid, cb, status='Enrolled', enrolled_on=date(2026, 9, 1))
    _add_fee(app, sid, 2000.0, date(2026, 9, 10))
    with app.app_context():
        m = course_wise_income_summary(date(2026, 9, 1), date(2026, 9, 30))
    assert m.get(cb) == 2000.0
    assert ca not in m


def test_course_wise_income_only_active_courses(app):
    ca = _mk_course(app, 'CWA2')
    cb = _mk_course(app, 'CWB2')
    sid = _mk_student(app, 3, [ca, cb])
    _set_enrollment(app, sid, ca, status='Completed', enrolled_on=date(2026, 8, 1),
                    completed_on=date(2026, 9, 15))
    _set_enrollment(app, sid, cb, status='Enrolled', enrolled_on=date(2026, 8, 1))
    _add_fee(app, sid, 2000.0, date(2026, 9, 10))
    _add_fee(app, sid, 2000.0, date(2026, 9, 20))
    with app.app_context():
        m = course_wise_income_summary(date(2026, 9, 1), date(2026, 9, 30))
    assert m.get(ca) == 1000.0
    assert m.get(cb) == 3000.0


def test_course_wise_totals_equal_collected(app):
    ca = _mk_course(app, 'RCA')
    cb = _mk_course(app, 'RCB')
    cc = _mk_course(app, 'RCC')
    sid = _mk_student(app, 4, [ca, cb, cc])
    _set_enrollment(app, sid, ca, status='Enrolled', enrolled_on=date(2026, 9, 1))
    _set_enrollment(app, sid, cb, status='Dropped', enrolled_on=date(2026, 8, 1),
                    completed_on=date(2026, 8, 20))
    _set_enrollment(app, sid, cc, status='Enrolled', enrolled_on=date(2026, 9, 5))
    _add_fee(app, sid, 1000.0, date(2026, 9, 8))
    _add_fee(app, sid, 2500.0, date(2026, 9, 12))
    with app.app_context():
        m = course_wise_income_summary(date(2026, 9, 1), date(2026, 9, 30))
    assert abs(sum(m.values()) - 3500.0) < 0.01
    assert cb not in m


# -- B2: company attribution falls back for companies without accounts ---------
def test_filter_by_company_methods_falls_back_to_company_id(app):
    cid = _mk_company(app, 'Fallback Co')
    _add_expense(app, 1000.0, date(2026, 9, 5), company_id=cid)
    with app.app_context():
        q = Expense.query.filter(Expense.expense_date >= date(2026, 9, 1),
                                 Expense.expense_date <= date(2026, 9, 30))
        rows = filter_by_company_methods(q, Expense.payment_method, cid).all()
    assert len(rows) == 1 and rows[0].amount == 1000.0


def test_company_without_accounts_reports_show_expenses(admin_client, app):
    cid = _mk_company(app, 'NoAcc Co')
    course_id = _mk_course(app, 'NOA', company_id=cid)
    sid = _mk_student(app, 5, [course_id])
    _set_enrollment(app, sid, course_id, status='Enrolled', enrolled_on=date.today())
    _add_fee(app, sid, 2000.0, date.today(), company_id=cid)
    _add_expense(app, 1000.0, date.today(), company_id=cid)

    body = admin_client.get(f'/reports?tab=expense&quick=today&company_id={cid}').get_data(as_text=True)
    assert '1,000.00' in body

    body_ov = admin_client.get(f'/reports?tab=overall&quick=today&company_id={cid}').get_data(as_text=True)
    assert '₹1,000.00' in body_ov


# -- B3: Excel "Monthly Avg" uses months elapsed in the year -------------------
def test_excel_monthly_avg_uses_elapsed_months(admin_client, app):
    today = date.today()
    with app.app_context():
        student = Student.query.filter_by(email='student@guha.test').first()
        db.session.add(FeeRecord(student_id=student.id, amount_paid=6000.0,
                                 taxable_amount=6000.0, gst_amount=0.0,
                                 payment_date=today, payment_method='Cash'))
        db.session.commit()
    resp = admin_client.get(f'/reports/excel?tab=income&month={today.month}&year={today.year}')
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    ws = wb['Summary']
    vals = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
    assert vals[f'Months elapsed in {today.year}'] == today.month
    assert abs(vals[f'Monthly Avg ({today.year})'] - 6000.0 / today.month) < 0.001


def test_excel_monthly_avg_full_year(admin_client, app):
    with app.app_context():
        student = Student.query.filter_by(email='student@guha.test').first()
        db.session.add(FeeRecord(student_id=student.id, amount_paid=1200.0,
                                 taxable_amount=1200.0, gst_amount=0.0,
                                 payment_date=date(2025, 6, 15), payment_method='Cash'))
        db.session.commit()
    resp = admin_client.get('/reports/excel?tab=income&year=2025&month=6')
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    ws = wb['Summary']
    vals = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
    assert vals['Months elapsed in 2025'] == 12
    assert abs(vals['Monthly Avg (2025)'] - 100.0) < 0.001


# -- B4: staff fees share the 30-day window with attendance --------------------
def test_staff_fees_collected_uses_30_day_window(staff_client, app):
    today = date.today()
    with app.app_context():
        student = Student.query.filter_by(email='student@guha.test').first()
        db.session.add(FeeRecord(student_id=student.id, amount_paid=1000.0,
                                 taxable_amount=1000.0, gst_amount=0.0,
                                 payment_date=today - timedelta(days=40), payment_method='Cash'))
        db.session.add(FeeRecord(student_id=student.id, amount_paid=500.0,
                                 taxable_amount=500.0, gst_amount=0.0,
                                 payment_date=today - timedelta(days=5), payment_method='Cash'))
        db.session.commit()
    body = staff_client.get('/reports').get_data(as_text=True)
    assert 'Fees Collected (30d)' in body
    assert '₹500.00' in body
    assert '₹1,000.00' not in body


def _add_account(app, name, company_id, acct_type='Cash'):
    with app.app_context():
        acc = Account.query.filter_by(name=name).first()
        if acc:
            acc.company_id = company_id
            acc.is_active = True
        else:
            db.session.add(Account(name=name, account_type=acct_type,
                                   company_id=company_id, is_active=True))
        db.session.commit()


# -- F1: unattributable fees surface as an "Unassigned" bucket -----------------
def test_course_wise_unassigned_for_student_without_enrollment(app):
    sid = _mk_student(app, 10, [])
    _add_fee(app, sid, 800.0, date.today())
    with app.app_context():
        m = course_wise_income_summary(date.today() - timedelta(days=1), date.today())
    assert None in m
    assert abs(m.get(None, 0.0) - 800.0) < 0.01


def test_course_wise_dropped_only_falls_back_not_unassigned(app):
    cid = _mk_course(app, 'DOP')
    sid = _mk_student(app, 11, [cid])
    _set_enrollment(app, sid, cid, status='Dropped', enrolled_on=date(2026, 8, 1),
                    completed_on=date(2026, 8, 15))
    _add_fee(app, sid, 900.0, date.today())
    with app.app_context():
        m = course_wise_income_summary(date.today() - timedelta(days=1), date.today())
    assert None not in m
    assert abs(m.get(cid, 0.0) - 900.0) < 0.01


def test_course_wise_unassigned_row_renders_on_web(admin_client, app):
    sid = _mk_student(app, 12, [])
    _add_fee(app, sid, 750.0, date.today())
    body = admin_client.get('/reports?tab=fees&quick=today').get_data(as_text=True)
    assert 'Unassigned' in body
    assert '750.00' in body


# -- F2: unified income attribution (direct company tag > unattributed) --------
def test_filter_fee_direct_company_wins_over_method(app):
    cid = _mk_company(app, 'Direct Co')
    _add_account(app, 'Cash', cid)
    course_id = _mk_course(app, 'DIR')
    sid = _mk_student(app, 13, [course_id])
    _set_enrollment(app, sid, course_id, status='Enrolled', enrolled_on=date(2026, 9, 1))
    _add_fee(app, sid, 1500.0, date(2026, 9, 10), method='UPI', company_id=cid)
    with app.app_context():
        q = FeeRecord.query.filter(FeeRecord.payment_date >= date(2026, 9, 1),
                                   FeeRecord.payment_date <= date(2026, 9, 30))
        rows = filter_by_company_methods(q, FeeRecord.payment_method, cid).all()
    assert len(rows) == 1 and rows[0].amount_paid == 1500.0


def test_filter_fee_unattributed_matched_included_direct_kept(app):
    cid = _mk_company(app, 'Match Co')
    _add_account(app, 'Cash', cid)
    course_id = _mk_course(app, 'MTH')
    sid = _mk_student(app, 14, [course_id])
    _set_enrollment(app, sid, course_id, status='Enrolled', enrolled_on=date(2026, 9, 1))
    _add_fee(app, sid, 1000.0, date(2026, 9, 10), method='Cash', company_id=None)
    _add_fee(app, sid, 2000.0, date(2026, 9, 11), method='UPI', company_id=cid)
    _add_fee(app, sid, 4000.0, date(2026, 9, 12), method='UPI', company_id=None)
    with app.app_context():
        q = FeeRecord.query.filter(FeeRecord.payment_date >= date(2026, 9, 1),
                                   FeeRecord.payment_date <= date(2026, 9, 30))
        rows = filter_by_company_methods(q, FeeRecord.payment_method, cid).all()
    assert sorted(r.amount_paid for r in rows) == [1000.0, 2000.0]


def test_course_wise_scope_attaches_unattributed_matched_fees(app):
    cid = _mk_company(app, 'Scope Co')
    _add_account(app, 'Cash', cid)
    course_id = _mk_course(app, 'SCP')
    sid = _mk_student(app, 15, [course_id])
    _set_enrollment(app, sid, course_id, status='Enrolled', enrolled_on=date.today())
    _add_fee(app, sid, 1200.0, date.today(), company_id=cid, method='Cash')
    _add_fee(app, sid, 600.0, date.today(), company_id=None, method='Cash')
    with app.app_context():
        m = course_wise_income_summary(date.today() - timedelta(days=1), date.today(),
                                       company_id=cid)
    assert abs(m.get(course_id, 0.0) - 1800.0) < 0.01


# -- F4: per-method export-all Excel (unbounded) -------------------------------
def test_payment_methods_export_all_excel(admin_client, app, monkeypatch):
    import app.routes.reports as reports_mod
    monkeypatch.setattr(reports_mod, 'PM_EXPORT_DETAIL_LIMIT', 2)
    sid = _mk_student(app, 16, [])
    for _ in range(5):
        _add_fee(app, sid, 100.0, date.today())
    resp = admin_client.get('/reports/excel?tab=payment_methods&quick=today&method=Cash')
    assert resp.status_code == 200
    assert 'Cash_Payments' in resp.headers.get('Content-Disposition', '')
    wb = load_workbook(BytesIO(resp.data))
    ws = wb.active
    assert ws.cell(row=2, column=1).value is not None
    assert ws.cell(row=8, column=1).value == 'Total'
    assert ws.cell(row=8, column=4).value == 500.0
    assert ws.max_row == 8


# -- F5: 30-day window reflected in labels -------------------------------------
def test_staff_recent_fees_label_30d(staff_client):
    body = staff_client.get('/reports').get_data(as_text=True)
    assert 'Recent Fee Payments (30d)' in body


# -- B5: payment-method detail bounded + truncation disclosed ------------------
def test_payment_breakdown_helper_truncates(app):
    sid = _mk_student(app, 6, [])
    for i in range(3):
        _add_fee(app, sid, 1000.0, date.today())
    with app.app_context():
        report, total = payment_method_period_breakdown(
            date.today() - timedelta(days=5), date.today(), detail_limit=2)
    assert report['Cash']['count'] == 3
    assert len(report['Cash']['records']) == 2
    assert report['Cash']['truncated'] == 1
    assert total == 3000.0


def test_payment_methods_web_shows_truncation_note(admin_client, app):
    sid = _mk_student(app, 7, [])
    for _ in range(51):
        _add_fee(app, sid, 100.0, date.today())
    body = admin_client.get('/reports?tab=payment_methods&quick=today').get_data(as_text=True)
    assert 'Showing latest 50 of 51 payments.' in body


def test_pdf_payment_methods_bounded(admin_client, app, monkeypatch):
    import app.routes.reports as reports_mod
    monkeypatch.setattr(reports_mod, 'PM_EXPORT_DETAIL_LIMIT', 2)
    sid = _mk_student(app, 8, [])
    for _ in range(3):
        _add_fee(app, sid, 100.0, date.today())
    resp = admin_client.get('/reports/pdf?tab=payment_methods&quick=today')
    assert resp.status_code == 200
    assert resp.mimetype == 'application/pdf'


def test_excel_payment_methods_bounded(admin_client, app, monkeypatch):
    import app.routes.reports as reports_mod
    monkeypatch.setattr(reports_mod, 'PM_EXPORT_DETAIL_LIMIT', 2)
    sid = _mk_student(app, 9, [])
    for _ in range(3):
        _add_fee(app, sid, 100.0, date.today())
    resp = admin_client.get('/reports/excel?tab=payment_methods&quick=today')
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    ws = wb['Cash']
    assert ws.max_row == 4
    assert ws.cell(row=4, column=1).value == 'Showing most recent 2 of 3 payments'
    summ = wb['Payment Summary']
    assert summ.cell(row=2, column=1).value == 'Cash'
    assert summ.cell(row=2, column=2).value == 300.0


# -- U1: inline warning when a selected company has no payment accounts --------
def test_no_accounts_warning_shown(app, admin_client):
    cid = _mk_company(app, 'NoAcc Warn Co')
    body = admin_client.get(f'/reports?tab=income&quick=today&company_id={cid}').get_data(as_text=True)
    assert 'has no payment accounts configured yet' in body


def test_no_accounts_warning_hidden_when_accounts_exist(app, admin_client):
    cid = _mk_company(app, 'HasAcc Co')
    _add_account(app, 'Cash', cid)
    body = admin_client.get(f'/reports?tab=income&quick=today&company_id={cid}').get_data(as_text=True)
    assert 'has no payment accounts configured yet' not in body


# -- U2: course-wise table explains the enrollment-window proration ------------
def test_course_wise_annotation_explains_window(app, admin_client):
    body = admin_client.get('/reports?tab=fees&quick=today').get_data(as_text=True)
    assert 'active on the payment date' in body
    assert 'Unassigned' in body


# -- U3: empty states are scoped to the selected company -----------------------
def test_empty_state_scoped_to_company(app, admin_client):
    cid = _mk_company(app, 'Empty State Co')
    body = admin_client.get(f'/reports?tab=income&quick=today&company_id={cid}').get_data(as_text=True)
    assert 'No income recorded in this period for Empty State Co' in body


def test_empty_state_not_scoped_without_company(app, admin_client):
    body = admin_client.get('/reports?tab=income&quick=today').get_data(as_text=True)
    assert 'No income recorded in this period for ' not in body


# -- U4: previous-period trend shows the actual prior span ---------------------
def test_previous_period_trend_shows_span(app, admin_client):
    body = admin_client.get('/reports?tab=income').get_data(as_text=True)
    assert 'vs last month (' in body